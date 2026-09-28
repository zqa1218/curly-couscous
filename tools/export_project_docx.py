#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把企划数据渲染成一份图文并茂的拍摄方案 docx，并转成 PDF。

用法：
    python tools/export_project_docx.py payload.json

payload.json 由本地服务组装，包含企划内容、算好的天文时间与图片绝对路径。
排版复用 build_shoot_plan.py 里的样式，保证两份文档看起来是一套东西。
"""

from __future__ import annotations

import json
import subprocess
import sys
import time
from pathlib import Path

# 让 stdout 固定用 UTF-8，避免中文路径在 Windows 控制台编码下变成乱码
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm  # noqa: E402

import build_shoot_plan as base  # noqa: E402


def add_images(doc, paths: list[str], width_cm: float = 5.2, per_row: int = 3) -> None:
    """按行插入若干张图，宽度统一，排不满一行也照常输出。"""
    existing = [Path(p) for p in paths if p and Path(p).exists()]
    if not existing:
        base.add_para(doc, "（这条灵感没有找到可用的参考图）", size=9, color=base.GRAY)
        return
    for start in range(0, len(existing), per_row):
        chunk = existing[start : start + per_row]
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = base.Pt(6)
        for index, path in enumerate(chunk):
            run = paragraph.add_run()
            run.add_picture(str(path), width=Cm(width_cm))
            if index < len(chunk) - 1:
                paragraph.add_run("    ")


def add_image(doc, path: str, width_cm: float = 13.0) -> None:
    target = Path(path)
    if not target.exists():
        base.add_para(doc, "（灯位图文件缺失）", size=9, color=base.GRAY)
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = base.Pt(8)
    paragraph.add_run().add_picture(str(target), width=Cm(width_cm))


def build(payload: dict, docx_path: Path) -> None:
    doc = Document()
    base.setup_document(doc)

    # 标题区
    title = doc.add_paragraph(style="Title")
    base.set_run_font(title.add_run(payload["name"]), size=23, bold=True, color=base.BLACK)
    title.paragraph_format.space_after = base.Pt(6)
    title.paragraph_format.line_spacing = 1.15
    subtitle = "  ·  ".join(
        [
            part
            for part in [
                "拍摄方案",
                f"{payload['plan'].get('city', '')} {payload['plan'].get('date', '')}".strip(),
                f"导出 {payload['exportedAt']}",
            ]
            if part
        ]
    )
    base.add_para(doc, subtitle, size=10.5, color=base.GRAY, space_after=14)

    brief = payload["brief"]
    plan = payload["plan"]
    research = payload.get("research", {})

    # 1 项目概览
    base.add_heading(doc, "1  项目概览", 1)

    def flatten(value: str) -> str:
        """把多行输入压成一行，去掉首尾多余的分隔符"""
        return "；".join(part.strip() for part in (value or "").splitlines() if part.strip())

    base.add_data_table(
        doc,
        ["项目", "内容", "项目", "内容"],
        [
            ["拍摄对象", brief.get("subject", ""), "道具", flatten(brief.get("props", ""))],
            ["拍摄场景", flatten(brief.get("shootingScene", "")) or plan.get("venue", ""), "风格关键词", "、".join(brief.get("styleKeywords", []))],
            ["焦段", brief.get("focalLength", ""), "光圈", brief.get("aperture", "")],
            ["拍摄手法", brief.get("shootingTechnique", ""), "前期效果", brief.get("onSetEffect", "")],
            ["后期效果", brief.get("postProduction", ""), "选定场地", plan.get("venue", "")],
        ],
        [3.0, 4.9, 3.0, 4.9],
        aligns=["l", "l", "l", "l"],
        row_height_cm=0.9,
    )

    # 2 时间与地点
    base.add_heading(doc, "2  时间与地点", 1)
    sun = payload.get("sun")
    if sun:
        base.add_data_table(
            doc,
            ["日出", "正午", "日落", "昼长"],
            [[sun.get("sunrise") or "—", sun.get("solarNoon") or "—", sun.get("sunset") or "—", sun.get("dayLength") or "—"]],
            [3.9, 3.9, 3.9, 4.1],
            aligns=["c", "c", "c", "c"],
            row_height_cm=0.9,
        )
        golden = sun.get("goldenMorning") and sun.get("goldenEvening")
        blue = sun.get("blueMorning") and sun.get("blueEvening")
        if golden or blue:
            rows = []
            if golden:
                rows.append(["黄金时刻", f"早 {sun['goldenMorning'][0]}–{sun['goldenMorning'][1]}", f"晚 {sun['goldenEvening'][0]}–{sun['goldenEvening'][1]}"])
            if blue:
                rows.append(["蓝调时刻", f"早 {sun['blueMorning'][0]}–{sun['blueMorning'][1]}", f"晚 {sun['blueEvening'][0]}–{sun['blueEvening'][1]}"])
            base.add_data_table(
                doc,
                ["时段", "早晨", "傍晚"],
                rows,
                [3.0, 6.4, 6.4],
                aligns=["l", "l", "l"],
                row_height_cm=0.9,
            )
        base.add_para(
            doc,
            f"以上时间为程序按 {payload['plan'].get('city', '')} 的经纬度本地计算（标准时区，不考虑夏令时）。",
            size=9,
            color=base.GRAY,
        )
    else:
        base.add_para(doc, "（还没有选定城市与日期，无法计算日出日落）", size=9.5, color=base.GRAY)

    if plan.get("timeWindow"):
        base.add_para(doc, f"选定拍摄时段：{plan['timeWindow']}", space_before=6)

    # 3 人物与主题
    base.add_heading(doc, "3  人物与主题", 1)
    if research.get("character"):
        base.add_para(doc, research["character"], line=1.6)
    if research.get("theme"):
        base.add_para(doc, f"主题定调：{research['theme']}", space_before=6, line=1.6)
    if not research.get("character") and not research.get("theme"):
        base.add_para(doc, "（还没有做人物调研与主题定调）", size=9.5, color=base.GRAY)

    # 4 灵感与来源
    ideas = payload.get("ideas", [])
    base.add_heading(doc, "4  灵感与来源", 1)
    if not ideas:
        base.add_para(doc, "（还没有定稿灵感）", size=9.5, color=base.GRAY)
    for idea in ideas:
        label = "完整想法" if idea["source"] == "idea" else f"组成元素组 · {len(idea['images'])} 张来源图"
        base.add_heading(doc, f"4.{idea['index']}  灵感 {idea['index']}（{label}）", 2)
        add_images(doc, idea["images"])
        if idea["elements"]:
            base.add_para(doc, "用到的元素：" + "、".join(idea["elements"]), size=9.5, color=base.GRAY, space_after=4)
        if idea.get("text"):
            base.add_para(doc, f"灵感正文：{idea['text']}", line=1.6)
        if idea.get("plan"):
            base.add_para(doc, idea["plan"], line=1.6, space_before=2)

    # 5 布光
    lighting = payload.get("lighting", [])
    base.add_heading(doc, "5  布光", 1)
    if not lighting:
        base.add_para(doc, "（还没有配布光方案）", size=9.5, color=base.GRAY)
    for entry in lighting:
        base.add_heading(doc, f"5.{entry['index']}  灵感 {entry['ideaIndex']} 的布光", 2)
        if entry.get("image"):
            add_image(doc, entry["image"])
            base.add_caption(doc, f"图 5-{entry['index']}  上传的灯位图")
        if entry.get("units"):
            base.add_data_table(
                doc,
                ["灯位", "灯具 功率 附件 位置 作用"],
                [[unit["title"], unit["detail"]] for unit in entry["units"]],
                [3.2, 12.6],
                aligns=["l", "l"],
                row_height_cm=1.1,
            )
        if entry.get("note"):
            base.add_para(doc, f"现场备注：{entry['note']}", size=9.5, color=base.GRAY)
        if not entry.get("image") and not entry.get("units"):
            base.add_para(doc, "（这条灵感还没有布光方案）", size=9.5, color=base.GRAY)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def docx_to_pdf(docx_path: Path, pdf_path: Path) -> bool:
    """用本机 Word 把 docx 转成 PDF；没有 Word 时返回 False，不影响 docx 交付。"""
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    script = (
        "$ErrorActionPreference='Stop';"
        "$w=New-Object -ComObject Word.Application;"
        "$w.Visible=$false;$w.DisplayAlerts=0;"
        f"$d=$w.Documents.Open('{docx_path}', $false, $true, $false);"
        f"$d.SaveAs2('{pdf_path}', 17);"
        "$d.Close();$w.Quit()"
    )
    try:
        result = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
            check=False,
            timeout=300,
            capture_output=True,
        )
        if result.returncode != 0:
            print(f"PowerShell 退出码 {result.returncode}", file=sys.stderr)
            print(result.stdout.decode("utf-8", "replace")[-600:], file=sys.stderr)
            print(result.stderr.decode("utf-8", "replace")[-600:], file=sys.stderr)
    except Exception as error:  # noqa: BLE001
        print(f"PDF 转换失败：{error}", file=sys.stderr)
        return False

    # Word 有时在进程退出之后才把文件写完，这里最多等 15 秒再确认
    for _ in range(30):
        if pdf_path.exists() and pdf_path.stat().st_size > 0:
            return True
        time.sleep(0.5)
    return False


def main() -> int:
    if len(sys.argv) < 2:
        print("用法：export_project_docx.py payload.json", file=sys.stderr)
        return 2
    payload = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    output_dir = Path(payload["outputDir"])
    stem = payload["fileStem"]
    docx_path = output_dir / f"{stem}.docx"
    pdf_path = output_dir / f"{stem}.pdf"

    build(payload, docx_path)
    to_pdf = payload.get("makePdf", True)
    pdf_ok = docx_to_pdf(docx_path, pdf_path) if to_pdf else False

    print(json.dumps({"docx": str(docx_path), "pdf": str(pdf_path) if pdf_ok else ""}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
