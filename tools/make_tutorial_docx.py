#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《印样台使用教程》docx，并转成 PDF。用法：python tools/make_tutorial_docx.py"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm  # noqa: E402

import build_shoot_plan as base  # noqa: E402
from export_project_docx import docx_to_pdf  # noqa: E402

REPO = Path(__file__).resolve().parent.parent
# 截图由 Chrome 无头模式输出：1440×900 纯页面区域，无浏览器外壳
QA = REPO / "docs" / "qa" / "final"


def picture(doc, name: str, caption: str, width_cm: float = 15.0) -> None:
    """按不放大原图的方式插图：截图 1440px 宽，15cm 对应约 244 DPI。"""
    path = QA / name
    if not path.exists():
        return
    try:
        from PIL import Image

        with Image.open(path) as probe:
            px_width = probe.width
        # 目标 DPI 下限 200：图片越窄，显示的物理尺寸越小，避免放大
        max_width_cm = px_width / 200 * 2.54
        width_cm = min(width_cm, max_width_cm)
    except Exception:  # noqa: BLE001
        pass
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = base.Pt(6)
    paragraph.paragraph_format.space_after = base.Pt(4)
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    base.add_caption(doc, caption)


def body(doc, text: str) -> None:
    base.add_para(doc, text, line=1.7, space_after=6)


def bullets(doc, items: list[str]) -> None:
    for item in items:
        paragraph = base.add_para(doc, f"· {item}", line=1.7, space_after=2)
        paragraph.paragraph_format.left_indent = Cm(0.5)


def build(docx_path: Path) -> None:
    doc = Document()
    base.setup_document(doc)

    title = doc.add_paragraph(style="Title")
    base.set_run_font(title.add_run("印样台使用教程"), size=23, bold=True, color=base.BLACK)
    title.paragraph_format.space_after = base.Pt(4)
    base.add_para(
        doc,
        "摄影企划工作台 · 从一句风格关键词到一份带布光的图文方案",
        size=10.5,
        color=base.GRAY,
        space_after=14,
    )

    base.add_heading(doc, "1  它是什么", 1)
    body(
        doc,
        "印样台把一次拍摄的全部准备工作收在一条流程里：先立项，再定时间地点与主题，然后找参考图、"
        "在黑板里拆解元素与灵感，最后每条灵感配一份布光，输出成图文方案。"
        "AI 在每一步先给候选，你点采纳之后才写进企划；所有灵感都能反查到来源图。",
    )
    body(
        doc,
        "它跑在你自己的电脑上：企划数据存在项目文件夹里，AI 密钥也只存在本机，不经过任何第三方服务。",
    )

    base.add_heading(doc, "2  快速开始", 1)
    bullets(
        doc,
        [
            "双击桌面上的「印样台」快捷方式，它会自动启动本地服务并打开浏览器。",
            "第一次打开会停在首页，点「进入项目」。",
            "在项目页点「新建项目」，起个名字，就进入工作台。",
            "如果还没有配置 AI：点右上角「AI 设置」，把 DeepSeek 的 API 密钥粘进去保存。",
            "没有密钥也能用，只是所有 AI 按钮会提示未配置，其余功能照常。",
        ],
    )
    picture(doc, "home.png", "图 1  首页")

    base.add_heading(doc, "3  一次拍摄的完整流程", 1)
    base.add_data_table(
        doc,
        ["阶段", "你要做的事", "AI 帮你做什么", "产出"],
        [
            ["S0 立项输入", "填人物、道具、场景、风格、焦段光圈、手法与前后期效果", "推荐道具、场景、风格、手法与效果，也能细化你写的文字", "项目卡"],
            ["S1 选址与时间", "选城市与日期，挑场地和时段", "推荐可落地的场地类型与光线条件", "日出日落、黄金与蓝调时刻（程序本地算）"],
            ["S2 人物与主题", "挑一个主题定调方向", "查人物或角色资料，给三个主题方向", "人物档案与主题"],
            ["S3 参考例图", "按目录去各站点找图，存进口袋", "给搜索词与直达链接（不做抓取）", "参考图库与来源索引"],
            ["S4 黑板", "翻口袋选图、勾元素、写灵感、连线分组", "拆元素、把一组元素整合成完整想法", "公共区布局与灵感"],
            ["S5 灵感归档", "同步、删改、定稿", "增量同步，不覆盖你的改动", "定稿灵感清单（带来源图）"],
            ["S6 布光", "上传自己的灯位图，或写要求让 AI 出方案", "按灯位逐条给灯具、功率、附件、位置与作用", "灯位表或灯位图"],
            ["S7 导出方案", "点一下导出", "—", "图文 PDF 与 docx"],
        ],
        [2.6, 4.6, 4.6, 4.0],
        size=9,
        row_height_cm=1.0,
    )

    base.add_heading(doc, "4  逐阶段说明", 1)

    base.add_heading(doc, "4.1  S0 立项输入", 2)
    body(doc, "分四组填写：拍摄对象、拍摄场景、风格与镜头、执行与效果。每个关键字段旁边都有 AI 按钮。")
    bullets(
        doc,
        [
            "道具、场景、风格、手法、前后期效果都可以让 AI 推荐，推荐结果默认全选，你可以取消不要的。",
            "「AI 细化我的写法」是把你已经写的内容改得更具体，选一条替换原文。",
            "焦段与光圈会一路带到后面的分析与灵感里，作为判断景深和透视的依据。",
            "场景一栏写「去哪拍」，手法一栏写「怎么拍」，S1 的选址建议只看场景栏。",
            "单行输入里按回车保存，任何位置按 Ctrl+S 也能保存。",
        ],
    )
    picture(doc, "s0.png", "图 2  S0 立项输入")

    base.add_heading(doc, "4.2  S1 选址与时间", 2)
    body(
        doc,
        "填城市和日期，程序立刻算出这一天的日出、正午、日落、昼长，以及早晨与傍晚的黄金时刻和蓝调时刻。"
        "这些时间由程序按经纬度本地计算，不是模型猜的。下方六个时段按钮点一下就把具体时间写进方案。",
    )
    body(doc, "场地交给 AI：它会给出场地类型、光线条件、最佳时段、需要提前确认的限制，以及和当前风格的匹配理由。")
    picture(doc, "s1.png", "图 3  S1 选址与时间")

    base.add_heading(doc, "4.3  S2 人物与主题", 2)
    body(
        doc,
        "AI 整理人物档案，覆盖出典身份、外形特征、服装配色、标志性元素、容易拍错的地方与可选气质。"
        "不确定的信息它会标明「需要核实」并说明该查什么，不编造设定。",
    )
    body(doc, "主题定调会给三个彼此有明显区别的方向，每个写清一句话主张、视觉走向和具体拍摄建议。")
    picture(doc, "s2.png", "图 4  S2 人物档案与主题定调")

    base.add_heading(doc, "4.4  S3 参考例图", 2)
    body(
        doc,
        "程序化抓取行不通——多数图站会拒绝程序访问——所以这里只做两件事：告诉你去哪找，帮你把找到的图分门别类存好。"
        "页面上方是按用途分好组的站点目录，填一个搜索词，所有站点的直达链接都会带上它。",
    )
    body(doc, "图存进三个口袋：抽象风格画作、优秀原画、优秀 Cos 作品。导入时记下来源链接、作者、站点与备注，重复的图会按内容自动跳过。")
    picture(doc, "s3.png", "图 5  S3 找图目录与口袋")

    base.add_heading(doc, "4.5  S4 黑板（核心）", 2)
    body(doc, "S4 从上到下四层：口袋翻页条、布置台、公共区、连线分组。")
    bullets(
        doc,
        [
            "顶端口袋翻页条：一页四张，左右翻页，点一张就送进布置台。",
            "布置台第一步：让 AI 看图，列出可以抽离复用的元素。",
            "第二步：勾选真正要用的元素，抽离成标签；写自己的灵感，也可以让 AI 补全成完整策划。",
            "第三步：决定以什么身份放进公共区——「完整想法」是终态不再参与连线；「组成元素」必须和别的元素连起来。",
            "公共区里的图可以任意拖动，卡片上标着身份与元素数量。",
            "连线把组成元素串成一组，可以连着连第三个、第四个；面板会显示「这一组（N 个组成元素）」，再由 AI 或你补全成完整想法。",
        ],
    )
    picture(doc, "s4.png", "图 6  S4 布置台与公共区")

    base.add_heading(doc, "4.6  S5 灵感归档", 2)
    body(
        doc,
        "点「从公共区同步」，把补全过的完整想法与分组结果收成一份清单。每条都带着来源图缩略图、用到的元素、"
        "灵感正文和完整策划，还能标作用环节（造型、光线、镜头、构图、道具、后期）。",
    )
    body(doc, "同步是增量的：你改过的文字、标好的环节、移出定稿的决定，重新同步都不会被覆盖。")
    picture(doc, "s5.png", "图 7  S5 灵感清单")

    base.add_heading(doc, "4.7  S6 布光", 2)
    body(doc, "布光按每条定稿灵感各配一份，两种方式可以任选或并用。")
    bullets(
        doc,
        [
            "上传灯位图：自己画的俯视图、现场照片、手绘扫描件都行。",
            "AI 推荐灯位：先写要求（只有一盏灯、要保留窗光层次、预算有限），它会按灯位逐条给出灯具类型、功率、附件、位置与高度、以及在这张画面里起什么作用；自然光和反光板也会单独列一条。",
            "每条灵感还留了现场备注，用来记「现场只有两个插座」这类实际限制。",
        ],
    )
    picture(doc, "s6.png", "图 8  S6 灯位表")

    base.add_heading(doc, "4.8  S7 导出方案", 2)
    body(
        doc,
        "点「导出方案」，约半分钟后在项目文件夹的 output 下生成 docx 与 PDF，页面上直接给下载链接。"
        "文档五章：项目概览、时间与地点（含黄金与蓝调时刻）、人物与主题、灵感与来源（每条灵感带来源图与元素）、布光。",
    )
    picture(doc, "s7.png", "图 9  S7 导出")

    base.add_heading(doc, "5  项目：新建、导入与蓝本", 1)
    body(doc, "每次拍摄是一个独立项目。项目页可以看到全部项目、新建，或者导入别人分享的项目包。")
    bullets(
        doc,
        [
            "新建项目：起个名字就开始。",
            "导入项目包：粘贴别人发给你的 JSON，或者直接选 .json 文件。项目包只带文字与设定，参考图需要自己重新收集。",
            "导出项目包：把当前项目导成一个 JSON 文件，可以发给别人。",
            "以它为蓝本：填新项目名、这次的拍摄对象、这次想要的风格，让 AI 给出三个调整方向（哪些保留、哪些必须改、具体怎么改），确认后创建。新项目会继承蓝本的场景骨架、主题、风格与布光文本，人物换成你填的这个，参考图与公共区不复制。",
        ],
    )
    picture(doc, "projects.png", "图 10  项目页")

    base.add_heading(doc, "6  数据放在哪", 1)
    body(doc, "每个项目是 projects 目录下的一个文件夹，拷走整个文件夹就带走了全部内容：")
    base.add_data_table(
        doc,
        ["目录", "内容"],
        [
            ["project.json", "立项、选址、人物档案、灵感清单、布光、公共区，全部结构化数据"],
            ["assets/refs", "收集进来的参考图原图"],
            ["assets/uploads", "你上传的灯位图等"],
            ["assets/thumbs", "列表用的缩略图"],
            ["board", "公共区画布"],
            ["output", "导出的 docx 与 PDF"],
        ],
        [4.2, 11.6],
        size=9.5,
        row_height_cm=0.9,
    )

    base.add_heading(doc, "7  常见问题", 1)
    bullets(
        doc,
        [
            "AI 按钮没反应或提示未配置：打开右上角「AI 设置」填密钥。",
            "提示密钥无效：确认复制完整，或去 DeepSeek 平台看余额。",
            "从网址导入例图失败：多数图站不允许程序下载，手动保存图片后再拖进来。",
            "导出的 PDF 没生成：docx 仍可用；PDF 依赖本机安装的 Word。",
            "想换台电脑继续：拷走整个项目文件夹，再在新机器上装好 Node 与依赖即可。",
        ],
    )

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main() -> int:
    out_dir = REPO / "docs"
    docx_path = out_dir / "印样台使用教程.docx"
    pdf_path = out_dir / "印样台使用教程.pdf"
    build(docx_path)
    ok = docx_to_pdf(docx_path, pdf_path)
    print(f"docx: {docx_path}")
    print(f"pdf : {pdf_path if ok else '（PDF 转换失败，docx 可用）'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
